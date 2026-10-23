"""Layer 0, step 1: parse course material into pages/slides (Build Plan 8.6).

Supports .md/.txt always; .pdf (pypdf or PyPDF2) and .pptx (python-pptx) when installed.
Headers, footers and page numbers are stripped; page/slide metadata is kept.
"""
from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass
from pathlib import Path


@dataclass
class Page:
    source: str
    page: int
    text: str
    chapter: str = ""


_PAGE_NO = re.compile(r"^\s*(page\s+)?\d+(\s*(/|of)\s*\d+)?\s*$", re.I)


def strip_boilerplate(pages: list[Page]) -> list[Page]:
    """Drop page-number lines and lines that repeat on most pages (headers/footers)."""
    if len(pages) >= 3:
        seen = Counter(line.strip() for p in pages for line in set(p.text.splitlines())
                       if line.strip())
        common = {line for line, n in seen.items() if n >= max(3, int(0.6 * len(pages)))}
    else:
        common = set()
    out = []
    for p in pages:
        keep = [ln for ln in p.text.splitlines()
                if not _PAGE_NO.match(ln) and ln.strip() not in common]
        out.append(Page(p.source, p.page, "\n".join(keep).strip(), p.chapter))
    return [p for p in out if p.text]


def _read_markdown(path: Path) -> list[Page]:
    """One page per top-level section; the heading is the chapter name."""
    text = path.read_text(encoding="utf-8", errors="replace")
    parts = re.split(r"(?m)^(?=# )", text)
    pages = []
    for i, part in enumerate(p for p in parts if p.strip()):
        m = re.match(r"# (.+)", part)
        pages.append(Page(path.name, i + 1, part.strip(), m.group(1).strip() if m else ""))
    return pages


def _read_pdf(path: Path) -> list[Page]:
    try:
        from pypdf import PdfReader
    except ImportError:
        try:
            from PyPDF2 import PdfReader  # type: ignore[no-redef]
        except ImportError as e:
            raise RuntimeError("reading PDFs needs `pip install pypdf` (or PyPDF2)") from e
    reader = PdfReader(str(path))
    return [Page(path.name, i + 1, (pg.extract_text() or "")) for i, pg in enumerate(reader.pages)]


def _read_pptx(path: Path) -> list[Page]:
    try:
        from pptx import Presentation
    except ImportError as e:
        raise RuntimeError("reading slides needs `pip install python-pptx`") from e
    pages = []
    for i, slide in enumerate(Presentation(str(path)).slides, start=1):
        texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]
        title = texts[0].splitlines()[0] if texts and texts[0] else ""
        pages.append(Page(path.name, i, "\n".join(texts), title))
    return pages


def load_documents(paths: list[str | Path]) -> list[Page]:
    """Load every supported file under the given files/directories."""
    files: list[Path] = []
    for p in map(Path, paths):
        files += sorted(f for f in p.rglob("*") if f.is_file()) if p.is_dir() else [p]
    pages: list[Page] = []
    readers = {".md": _read_markdown, ".txt": _read_markdown, ".pdf": _read_pdf,
               ".pptx": _read_pptx}
    for f in files:
        reader = readers.get(f.suffix.lower())
        if reader:
            pages += strip_boilerplate(reader(f))
    return pages
